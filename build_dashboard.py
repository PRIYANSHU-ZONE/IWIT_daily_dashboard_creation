#!/usr/bin/env python3
"""Build a standalone daily STN operations dashboard from CSV or Excel extracts."""

import argparse
import csv
import html
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path


SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xlsm"}
MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def newest_extract(data_dir, prefix):
    candidates = [
        path
        for path in data_dir.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_SUFFIXES
        and path.stem.lower().startswith(prefix.lower())
        and not path.name.startswith("~$")
    ]
    if not candidates:
        return None

    def rank(path):
        match = re.search(r"(?<!\d)(\d{1,2})[ _-]+([A-Za-z]{3})(?:[ _-]+(\d{2,4}))?", path.stem)
        if match:
            day, month, year = match.groups()
            year = int(year) if year else date.today().year
            if year < 100:
                year += 2000
            try:
                return (date(year, MONTHS.index(month.title()) + 1, int(day)), path.stat().st_mtime)
            except (ValueError, IndexError):
                pass
        return (date.min, path.stat().st_mtime)

    return max(candidates, key=rank)


def records_from(path):
    """Yield row dictionaries while keeping large CSV/XLSX inputs streamed."""
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as source:
            for row in csv.DictReader(source):
                yield row
        return

    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError("Excel input requires openpyxl. Install it with: pip install -r requirements.txt") from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows, ())
        header_names = [str(value).strip() if value is not None else "" for value in headers]
        for values in rows:
            yield dict(zip(header_names, values))
    finally:
        workbook.close()


def clean(value):
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.upper() in {"", "NA", "N/A", "NULL", "NONE"} else text


def quantity(value):
    text = clean(value).replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except (ValueError, TypeError):
        return 0.0


def compact(value):
    value = float(value)
    if abs(value) >= 1_000_000:
        return f"{value / 1_000_000:.1f}M"
    if abs(value) >= 1_000:
        return f"{value / 1_000:.1f}K"
    return f"{value:,.0f}"


def top_counts(groups, limit=10):
    return [
        {"label": label, "value": len(ids)}
        for label, ids in sorted(groups.items(), key=lambda pair: (-len(pair[1]), pair[0].casefold()))[:limit]
    ]


def build_data(stn_path, str_path):
    stn_ids = set()
    item_ids = set()
    stn_statuses = defaultdict(set)
    category_ids = {
        "vertical": defaultdict(set),
        "source": defaultdict(set),
        "movement": defaultdict(set),
    }
    status_ids = defaultdict(set)
    totals = Counter()

    for row in records_from(stn_path):
        totals["item_rows"] += 1
        stn_id = clean(row.get("STN ID"))
        item_id = clean(row.get("STN Item ID"))
        if stn_id:
            stn_ids.add(stn_id)
            status = clean(row.get("Status")).upper() or "UNKNOWN"
            status_ids[status].add(stn_id)
            stn_statuses[stn_id].add(status)
            for key, field in (("vertical", "Vertical"), ("source", "Source Warehouse"), ("movement", "Movement Type")):
                category_ids[key][clean(row.get(field)) or "Unknown"].add(stn_id)
        if item_id:
            item_ids.add(item_id)
        totals["requested"] += quantity(row.get("Requested Quantity"))
        totals["scheduled"] += quantity(row.get("Scheduled Quantity"))

    str_user_ids = defaultdict(set)
    stn_user_status_ids = defaultdict(lambda: defaultdict(set))
    matched_str_ids = set()
    if str_path:
        for row in records_from(str_path):
            str_id = clean(row.get("STR ID"))
            user = clean(row.get("User")) or "Unknown"
            if str_id:
                str_user_ids[user].add(str_id)
                if str_id in stn_ids:
                    matched_str_ids.add(str_id)
                    for status in stn_statuses[str_id]:
                        stn_user_status_ids[user][status].add(str_id)

    scheduled_pct = totals["scheduled"] / totals["requested"] * 100 if totals["requested"] else 0
    return {
        "cards": [
            {"label": "Total STN IDs", "value": f"{len(stn_ids):,}"},
            {"label": "Total STN Items", "value": compact(len(item_ids) or totals["item_rows"])},
            {"label": "Requested Qty", "value": compact(totals["requested"])},
            {"label": "Scheduled Qty", "value": compact(totals["scheduled"])},
            {"label": "Scheduled Fulfilment", "value": f"{scheduled_pct:.2f}%"},
        ],
        "verticals": top_counts(category_ids["vertical"]),
        "sources": top_counts(category_ids["source"]),
        "movements": top_counts(category_ids["movement"], 20),
        "statuses": [
            {"label": label.title(), "value": len(ids)}
            for label, ids in sorted(status_ids.items(), key=lambda pair: pair[0])
        ],
        "strUsers": [
            {"label": label, "value": len(ids)}
            for label, ids in sorted(str_user_ids.items(), key=lambda pair: (-len(pair[1]), pair[0].casefold()))[:10]
        ],
        "stnUserMatrix": [
            {
                "user": user,
                "cancelled": len(statuses.get("CANCELLED", set())),
                "closed": len(statuses.get("CLOSED", set())),
                "lost": len(statuses.get("LOST", set())),
                "scheduled": len(statuses.get("SCHEDULED", set())),
                "total": len(set().union(*statuses.values())) if statuses else 0,
            }
            for user, statuses in sorted(
                stn_user_status_ids.items(),
                key=lambda pair: (-len(set().union(*pair[1].values())), pair[0].casefold()),
            )
        ],
        "hasStrFile": str_path is not None,
        "strStnOverlap": len(matched_str_ids),
        "itemRows": totals["item_rows"],
    }


def render_html(data, stn_path, str_path):
    payload = json.dumps(data, ensure_ascii=True, separators=(",", ":")).replace("</", "<\\/")
    input_files = [stn_path.name] + ([str_path.name] if str_path else [])
    file_list = html.escape(" | ".join(input_files))
    return f'''<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Daily STN Operations</title>
<style>
:root {{ color-scheme: light; --ink: #202b27; --muted: #68736e; --paper: #f4f6f1; --panel: #fff; --line: #dfe5de; --green: #247557; --lime: #c7dd6c; --rust: #c86a4b; --blue: #5f8da4; --gold: #d3a844; }}
* {{ box-sizing: border-box; }}
body {{ margin: 0; background: var(--paper); color: var(--ink); font: 14px/1.45 ui-sans-serif, system-ui, -apple-system, "Segoe UI", sans-serif; }}
main {{ max-width: 1220px; margin: 0 auto; padding: 34px 28px 46px; }}
header {{ display: flex; justify-content: space-between; align-items: end; gap: 24px; padding: 0 0 24px; border-bottom: 1px solid var(--line); }}
.eyebrow {{ color: var(--green); font-size: 11px; font-weight: 750; letter-spacing: .1em; text-transform: uppercase; }}
h1 {{ margin: 5px 0 0; font: 500 32px/1.05 Georgia, "Times New Roman", serif; }}
.stamp {{ color: var(--muted); font-size: 12px; text-align: right; }}
.stamp strong {{ display: block; color: var(--ink); font-size: 13px; font-weight: 650; }}
.cards {{ display: grid; grid-template-columns: repeat(5, minmax(0, 1fr)); gap: 10px; padding: 18px 0 24px; }}
.metric {{ min-height: 100px; background: var(--panel); border: 1px solid var(--line); border-top: 3px solid var(--green); padding: 15px 16px; }}
.metric:nth-child(2) {{ border-top-color: var(--blue); }} .metric:nth-child(3) {{ border-top-color: var(--gold); }} .metric:nth-child(4) {{ border-top-color: var(--rust); }} .metric:nth-child(5) {{ border-top-color: #91ad3e; }}
.metric-label {{ color: var(--muted); font-size: 12px; }} .metric-value {{ display: block; margin-top: 7px; font: 500 27px/1 Georgia, "Times New Roman", serif; font-variant-numeric: tabular-nums; }}
.grid {{ display: grid; grid-template-columns: minmax(0, 1fr) minmax(0, 1fr); gap: 12px; }}
.panel {{ background: var(--panel); border: 1px solid var(--line); padding: 17px 18px 18px; min-width: 0; }}
.panel h2 {{ margin: 0; font: 500 19px/1.2 Georgia, "Times New Roman", serif; }}
.sub {{ margin: 4px 0 14px; color: var(--muted); font-size: 11px; }}
.bar-list {{ display: grid; gap: 9px; }}
.bar-row {{ display: grid; grid-template-columns: minmax(100px, 155px) minmax(60px, 1fr) 48px; align-items: center; gap: 9px; min-height: 16px; }}
.bar-label {{ overflow: hidden; color: #43504a; font-size: 11px; text-overflow: ellipsis; white-space: nowrap; }}
.track {{ height: 9px; overflow: hidden; background: #edf0eb; }} .fill {{ height: 100%; background: var(--green); }}
.bar-value {{ color: var(--ink); font-size: 11px; font-variant-numeric: tabular-nums; text-align: right; }}
.lower {{ margin-top: 12px; }}
.status-layout {{ display: grid; grid-template-columns: minmax(130px, .8fr) minmax(130px, 1fr); align-items: center; gap: 12px; min-height: 190px; }}
.donut {{ display: block; width: min(100%, 190px); margin: auto; }} .donut-center {{ font: 500 18px Georgia, "Times New Roman", serif; fill: var(--ink); }} .donut-caption {{ font: 10px ui-sans-serif, system-ui, sans-serif; fill: var(--muted); }}
.legend {{ display: grid; gap: 9px; }} .legend-item {{ display: grid; grid-template-columns: 9px minmax(0, 1fr) auto; gap: 8px; align-items: center; font-size: 11px; }} .swatch {{ width: 9px; height: 9px; }} .legend-label {{ color: #43504a; }} .legend-value {{ font-variant-numeric: tabular-nums; }}
.table-wrap {{ overflow-x: auto; }} table {{ width: 100%; border-collapse: collapse; font-size: 11px; }} th, td {{ padding: 8px 10px; border-bottom: 1px solid var(--line); text-align: left; }} th {{ color: var(--muted); font-weight: 650; }} td:last-child, th:last-child {{ text-align: right; font-variant-numeric: tabular-nums; }}
.notice {{ margin-top: 12px; padding: 12px 14px; border-left: 3px solid var(--gold); background: #fff9e9; color: #594a25; font-size: 12px; }}
footer {{ display: flex; flex-wrap: wrap; justify-content: space-between; gap: 8px 20px; border-top: 1px solid var(--line); margin-top: 16px; padding-top: 12px; color: var(--muted); font-size: 10px; }}
@media (max-width: 760px) {{ main {{ padding: 22px 14px 30px; }} header {{ align-items: start; flex-direction: column; gap: 8px; }} .stamp {{ text-align: left; }} .cards {{ grid-template-columns: repeat(2, minmax(0, 1fr)); }} .metric:last-child {{ grid-column: span 2; }} .grid {{ grid-template-columns: 1fr; }} }}
@media (max-width: 420px) {{ h1 {{ font-size: 27px; }} .bar-row {{ grid-template-columns: minmax(84px, 115px) minmax(35px, 1fr) 42px; gap: 6px; }} .status-layout {{ grid-template-columns: 1fr 1fr; gap: 6px; }} }}
@media print {{ body {{ background: #fff; }} main {{ max-width: none; padding: 0; }} .panel, .metric {{ break-inside: avoid; }} }}
</style>
</head>
<body>
<main>
<header><div><div class="eyebrow">Supply chain | Daily operations</div><h1>STN movement pulse</h1></div><div class="stamp"><strong id="generated"></strong><span>Source extracts: {file_list}</span></div></header>
<section class="cards" id="cards" aria-label="Daily key performance indicators"></section>
<section class="grid">
  <article class="panel"><h2>Top 10 Verticals by STNs</h2><p class="sub">Distinct STN IDs per vertical</p><div class="bar-list" id="verticals"></div></article>
  <article class="panel"><h2>Top Source Warehouses</h2><p class="sub">Distinct STN IDs by source warehouse</p><div class="bar-list" id="sources"></div></article>
  <article class="panel"><h2>STNs by Movement Type</h2><p class="sub">Distinct STN IDs by movement type</p><div class="bar-list" id="movements"></div></article>
  <article class="panel"><h2>STN Status Distribution</h2><p class="sub">Distinct STN IDs by current status</p><div class="status-layout"><svg class="donut" viewBox="0 0 200 200" role="img" aria-label="STN status distribution"><g id="donut-slices" transform="rotate(-90 100 100)"></g><text id="donut-total" class="donut-center" x="100" y="97" text-anchor="middle"></text><text class="donut-caption" x="100" y="116" text-anchor="middle">STN IDs</text></svg><div class="legend" id="status-legend"></div></div></article>
    <article class="panel"><h2>STR IDs by User</h2><p class="sub" id="user-caption"></p><div class="bar-list" id="users"></div></article>
  <article class="panel"><h2>STN Status by User</h2><p class="sub">Distinct STN IDs grouped by creator and status</p><div class="table-wrap" id="user-matrix"></div></article>
</section>
<div class="notice" id="data-note"></div>
<footer><span>Quantities are summed from STN item rows; IDs are distinct-counted.</span><span>Standalone report · no external assets or network calls</span></footer>
</main>
<script>
const dashboardData = {payload};
const palette = ["#247557", "#c86a4b", "#5f8da4", "#d3a844", "#91ad3e", "#8b7660"];
document.getElementById("generated").textContent = "Generated " + new Date().toLocaleString();
const cards = document.getElementById("cards");
dashboardData.cards.forEach(item => {{
  const node = document.createElement("div"); node.className = "metric";
  const label = document.createElement("span"); label.className = "metric-label"; label.textContent = item.label;
  const value = document.createElement("strong"); value.className = "metric-value"; value.textContent = item.value;
  node.append(label, value); cards.append(node);
}});
function renderBars(targetId, rows) {{
  const target = document.getElementById(targetId); const max = Math.max(1, ...rows.map(row => row.value));
  if (!rows.length) {{ target.textContent = "No data available"; return; }}
  rows.forEach(row => {{
    const line = document.createElement("div"); line.className = "bar-row";
    const label = document.createElement("span"); label.className = "bar-label"; label.textContent = row.label; label.title = row.label;
    const track = document.createElement("div"); track.className = "track";
    const fill = document.createElement("div"); fill.className = "fill"; fill.style.width = (row.value / max * 100) + "%"; track.append(fill);
    const value = document.createElement("span"); value.className = "bar-value"; value.textContent = row.value.toLocaleString();
    line.append(label, track, value); target.append(line);
  }});
}}
renderBars("verticals", dashboardData.verticals);
renderBars("sources", dashboardData.sources);
renderBars("movements", dashboardData.movements);
renderBars("users", dashboardData.strUsers);
document.getElementById("user-caption").textContent = dashboardData.hasStrFile ? "Distinct STR IDs per user (STR extract)" : "STR extract not found";
const total = dashboardData.statuses.reduce((sum, row) => sum + row.value, 0);
document.getElementById("donut-total").textContent = total.toLocaleString();
const slices = document.getElementById("donut-slices"); const legend = document.getElementById("status-legend");
let offset = 0; const circumference = 2 * Math.PI * 70;
dashboardData.statuses.forEach((row, index) => {{
  const length = total ? row.value / total * circumference : 0;
  const circle = document.createElementNS("http://www.w3.org/2000/svg", "circle");
  circle.setAttribute("cx", "100"); circle.setAttribute("cy", "100"); circle.setAttribute("r", "70");
  circle.setAttribute("fill", "none"); circle.setAttribute("stroke", palette[index % palette.length]); circle.setAttribute("stroke-width", "25");
  circle.setAttribute("stroke-dasharray", `${{length}} ${{circumference - length}}`); circle.setAttribute("stroke-dashoffset", `${{-offset}}`); slices.append(circle); offset += length;
  const item = document.createElement("div"); item.className = "legend-item";
  const swatch = document.createElement("span"); swatch.className = "swatch"; swatch.style.background = palette[index % palette.length];
  const label = document.createElement("span"); label.className = "legend-label"; label.textContent = row.label;
  const value = document.createElement("strong"); value.className = "legend-value"; value.textContent = row.value.toLocaleString();
  item.append(swatch, label, value); legend.append(item);
}});
const matrix = document.getElementById("user-matrix");
if (!dashboardData.stnUserMatrix.length) {{
  matrix.textContent = "Not available: the supplied extracts do not share matching STN ID / STR ID values.";
}} else {{
  const table = document.createElement("table"); const head = document.createElement("thead"); const headRow = document.createElement("tr");
    ["User", "Cancelled", "Closed", "Lost", "Scheduled", "Total"].forEach(text => {{ const cell = document.createElement("th"); cell.textContent = text; headRow.append(cell); }});
  head.append(headRow); table.append(head); const body = document.createElement("tbody");
    dashboardData.stnUserMatrix.forEach(row => {{ const tr = document.createElement("tr"); [row.user, row.cancelled, row.closed, row.lost, row.scheduled, row.total].forEach((text, index) => {{ const cell = document.createElement("td"); cell.textContent = index ? text.toLocaleString() : text; tr.append(cell); }}); body.append(tr); }});
  table.append(body); matrix.append(table);
}}
const note = document.getElementById("data-note");
note.textContent = dashboardData.hasStrFile
  ? `Creator breakdown uses STR IDs. ${{dashboardData.strStnOverlap.toLocaleString()}} STR IDs matched an STN ID in these extracts; status-by-user is ${{dashboardData.stnUserMatrix.length ? "shown for matches" : "unavailable without a shared identifier"}}.`
  : "Only an STN extract was found. Creator charts require a matching STR extract; status-by-user also requires a shared identifier.";
</script>
</body>
</html>'''


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=Path(__file__).parent / "This_folder_has_daily_data")
    parser.add_argument("--output", type=Path, default=Path(__file__).parent / "daily_dashboard.html")
    args = parser.parse_args()
    if not args.data_dir.is_dir():
        parser.error(f"data directory does not exist: {args.data_dir}")

    stn_path = newest_extract(args.data_dir, "STN ID")
    if not stn_path:
        parser.error(f"no STN ID CSV/XLSX/XLSM extract found in {args.data_dir}")
    str_path = newest_extract(args.data_dir, "STR ID")
    try:
        data = build_data(stn_path, str_path)
    except RuntimeError as error:
        print(error, file=sys.stderr)
        return 1

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(render_html(data, stn_path, str_path), encoding="utf-8")
    print(f"Dashboard written to {args.output}")
    print(f"STN IDs: {data['cards'][0]['value']} | STN items: {data['cards'][1]['value']} | fulfilment: {data['cards'][4]['value']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())